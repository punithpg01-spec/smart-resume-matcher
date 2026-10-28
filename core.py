"""Core functions for Smart Resume Matcher, separated for testing."""
import csv
import io
from datetime import datetime

from docx import Document
from pypdf import PdfReader
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

from matcher import find_skills, recommendations

def extract_text(upload):
    """Extract text from an uploaded TXT, PDF, or DOCX file."""
    if upload is None:
        return ""

    filename = upload.name.lower()
    upload.seek(0)

    if filename.endswith(".txt"):
        return upload.getvalue().decode("utf-8", errors="replace")

    if filename.endswith(".pdf"):
        reader = PdfReader(upload)
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if filename.endswith(".docx"):
        document = Document(upload)
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_text = [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
        return "\n".join(paragraphs + table_text)

    raise ValueError("Unsupported file type. Please use PDF, DOCX, or TXT.")

def analyze(resume, job, resume_name, job_name):
    """Compare resume/job text and return a consistent results dictionary."""
    resume = resume.strip()
    job = job.strip()

    if not resume:
        raise ValueError("Resume text is empty.")
    if not job:
        raise ValueError("Job description is empty.")

    # TF-IDF requires at least some terms that are not removed as stop words.
    try:
        vectorizer = TfidfVectorizer(stop_words="english")
        vectors = vectorizer.fit_transform([resume, job])
        score = float(
            cosine_similarity(vectors[0:1], vectors[1:2])[0][0] * 100
        )
    except ValueError:
        score = 0.0

    resume_skills = set(find_skills(resume))
    job_skills = set(find_skills(job))
    matched = sorted(resume_skills & job_skills)
    missing = sorted(job_skills - resume_skills)

    # `recommendations` is a dictionary, so retrieve each item with .get().
    recommendation_lines = []
    for skill in missing:
        advice = recommendations.get(
            skill, "Build a project demonstrating " + skill + "."
        )
        recommendation_lines.append(f"{skill}: {advice}")

    return {
        "resume_name": resume_name,
        "job_name": job_name,
        "score": round(score, 2),
        "matched": matched,
        "missing": missing,
        "recommendations": recommendation_lines,
        "resume_skills": sorted(resume_skills),
        "job_skills": sorted(job_skills),
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }

def txt_report(result):
    """Create a downloadable plain-text report."""
    lines = [
        "SMART RESUME MATCHER REPORT",
        "=" * 30,
        f"Resume: {result['resume_name']}",
        f"Job description: {result['job_name']}",
        f"Date: {result['created_at']}",
        f"Text similarity: {result['score']}%",
        "",
        "MATCHED SKILLS",
        *[f"- {skill}" for skill in result["matched"]],
        "",
        "SKILLS TO IMPROVE",
        *[f"- {skill}" for skill in result["missing"]],
        "",
        "RECOMMENDATIONS",
        *[f"- {item}" for item in result["recommendations"]],
    ]
    return "\n".join(lines)

def csv_report(result):
    """Create a downloadable CSV report."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Category", "Value"])
    writer.writerow(["Resume", result["resume_name"]])
    writer.writerow(["Job description", result["job_name"]])
    writer.writerow(["Date", result["created_at"]])
    writer.writerow(["Text similarity (%)", result["score"]])

    for skill in result["matched"]:
        writer.writerow(["Matched skill", skill])
    for skill in result["missing"]:
        writer.writerow(["Skill to improve", skill])
    for item in result["recommendations"]:
        writer.writerow(["Recommendation", item])

    return output.getvalue()

def pdf_report(result):
    """Create a downloadable PDF report with basic pagination."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    output = io.BytesIO()
    pdf = canvas.Canvas(output, pagesize=letter)
    width, height = letter
    margin = 50
    y = height - margin

    def write_line(text, bold=False):
        nonlocal y
        if y < margin:
            pdf.showPage()
            y = height - margin
        pdf.setFont("Helvetica-Bold" if bold else "Helvetica", 10)
        # Basic PDF Helvetica fonts cannot render every Unicode character.
        safe_text = str(text).encode("latin-1", errors="replace").decode("latin-1")
        pdf.drawString(margin, y, safe_text[:110])
        y -= 17

    write_line("SMART RESUME MATCHER REPORT", True)
    write_line(f"Resume: {result['resume_name']}")
    write_line(f"Job description: {result['job_name']}")
    write_line(f"Date: {result['created_at']}")
    write_line(f"Text similarity: {result['score']}%")
    y -= 10

    write_line("MATCHED SKILLS", True)
    for skill in result["matched"] or ["No matching skills detected"]:
        write_line(f"- {skill}")

    y -= 10
    write_line("SKILLS TO IMPROVE", True)
    for skill in result["missing"] or ["No missing skills detected"]:
        write_line(f"- {skill}")

    y -= 10
    write_line("RECOMMENDATIONS", True)
    for item in result["recommendations"] or ["No recommendations available"]:
        write_line(f"- {item}")

    pdf.save()
    output.seek(0)
    return output.getvalue()
